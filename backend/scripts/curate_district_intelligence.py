"""Curate aggregate NISR SAS workbooks into the Module 8 runtime dataset.

This script intentionally imports only published aggregate tables. It never reads
microdata or farmer-level records. Run from backend/ after replacing source files.
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
SOURCE_DIR = ROOT / "data" / "nisr-audit"
OUTPUT = ROOT / "backend" / "app" / "data" / "nisr_district_intelligence.json"
BOUNDARIES = ROOT / "backend" / "app" / "data" / "rwanda_district_boundaries_2022.geojson"

FILES = {
    "A": SOURCE_DIR / "Tables_2026_Season_A.xlsx",
    "B": SOURCE_DIR / "Tables_2026_Season_B.xlsx",
}

SOURCE_META = {
    "A": {
        "title": "NISR Seasonal Agricultural Survey, Season A 2026 — Statistical Tables",
        "url": "https://www.statistics.gov.rw/sites/default/files/documents/2026-04/Tables_2026%20Season%20A.xlsx",
        "report_url": "https://www.gis.statistics.gov.rw/statistical-publications/agriculture-environment/seasonal-agricultural-survey-season-2026",
        "observation_period": "September 2025–February 2026",
        "sample_note": "NISR reports 1,200 sampled segments and 554 large-scale farmers across all 30 districts.",
    },
    "B": {
        "title": "NISR Seasonal Agricultural Survey, Season B 2026 — Statistical Tables",
        "url": "https://www.gis.statistics.gov.rw/sites/default/files/documents/2026-09/Tables_2026%20Season%20B.xlsx",
        "report_url": "https://www.gis.statistics.gov.rw/statistical-publications/agriculture-environment/seasonal-agricultural-survey-season-b-2026",
        "observation_period": "March–June 2026",
        "sample_note": "NISR reports 1,200 sampled segments and 505 large-scale farmers across all 30 districts.",
    },
}

SKIP_CROP_GROUPS = {
    "Cereals", "Tubers and Roots", "Banana", "Bananas", "Legumes and Pulses",
    "Bean", "Vegetables and Fruits", "Total Developed land",
}

INDICATORS = {
    "cultivated_area": {"label": "Cultivated area", "unit": "ha", "category": "crop"},
    "harvested_area": {"label": "Harvested area", "unit": "ha", "category": "crop"},
    "production": {"label": "Crop production", "unit": "metric tonnes", "category": "crop"},
    "yield": {"label": "Yield per hectare", "unit": "metric tonnes/ha", "category": "crop"},
    "agricultural_land": {"label": "Agricultural land", "unit": "ha", "category": "land"},
    "physical_cultivated_land": {"label": "Physical cultivated land", "unit": "ha", "category": "land"},
    "seasonal_crop_land": {"label": "Land under seasonal crops", "unit": "ha", "category": "land"},
    "modern_irrigated_area": {"label": "Modern irrigated agricultural land", "unit": "ha", "category": "irrigation"},
    "erosion_control_area": {"label": "Agricultural area under erosion control", "unit": "ha", "category": "practice"},
    "agroforestry_area": {"label": "Agricultural area under agroforestry", "unit": "ha", "category": "practice"},
    "inorganic_fertilizer_area": {"label": "Area using inorganic fertilizer", "unit": "ha", "category": "practice"},
    "organic_fertilizer_area": {"label": "Area using organic fertilizer", "unit": "ha", "category": "practice"},
    "erosion_control_farmers": {"label": "Farmers protecting land against erosion", "unit": "% of farmers", "category": "practice"},
    "mechanized_farmers": {"label": "Farmers using mechanical equipment", "unit": "% of farmers", "category": "practice"},
    "irrigating_farmers": {"label": "Farmers practicing irrigation", "unit": "% of farmers", "category": "irrigation"},
    "agroforestry_farmers": {"label": "Farmers practicing agroforestry", "unit": "% of farmers", "category": "practice"},
}


def clean(value):
    return " ".join(str(value).replace("\n", " ").split()).strip() if value is not None else ""


def numeric(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def workbook_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def province_map() -> dict[str, str]:
    geojson = json.loads(BOUNDARIES.read_text(encoding="utf-8"))
    return {clean(f["properties"]["district"]): clean(f["properties"]["province"]) for f in geojson["features"]}


def append_record(records, *, year, season, province, district, level, indicator, value, crop=None, table):
    if not numeric(value):
        return
    records.append({
        "year": int(year), "season": season, "province": province, "district": district,
        "geography_level": level, "crop": crop, "indicator": indicator,
        "value": round(float(value), 6), "unit": INDICATORS[indicator]["unit"],
        "source_season": season, "source_table": table,
    })


def parse_national(ws, season, records):
    # Published table contains current and comparison years in merged indicator groups.
    indicator = ""
    for col in range(3, ws.max_column + 1):
        if ws.cell(3, col).value:
            heading = clean(ws.cell(3, col).value).lower()
            indicator = (
                "cultivated_area" if "cultivated" in heading else
                "harvested_area" if "harvested" in heading else
                "production" if "production" in heading else
                "yield" if "yield" in heading else ""
            )
        period = clean(ws.cell(4, col).value)
        if not indicator or not period:
            continue
        year = int(period.split()[0])
        for row in range(5, ws.max_row + 1):
            crop = clean(ws.cell(row, 2).value)
            value = ws.cell(row, col).value
            if crop and crop not in SKIP_CROP_GROUPS:
                append_record(records, year=year, season=season, province=None, district=None, level="national", indicator=indicator, value=value, crop=crop, table="Table 6")


def parse_crop_table(ws, season, indicator, header_row, data_row, provinces, records, table):
    for row in range(data_row, ws.max_row + 1):
        district = clean(ws.cell(row, 2).value)
        if district not in provinces:
            continue
        for col in range(3, ws.max_column + 1):
            crop = clean(ws.cell(header_row, col).value)
            if not crop or crop in SKIP_CROP_GROUPS:
                continue
            value = ws.cell(row, col).value
            if indicator == "yield" and numeric(value):
                value = value / 1000.0  # Published district table is kg/ha; runtime unit is t/ha.
            append_record(records, year=2026, season=season, province=provinces[district], district=district, level="district", indicator=indicator, value=value, crop=crop, table=table)


def parse_land(ws, season, provinces, records, table):
    mapping = {4: "agricultural_land", 7: "physical_cultivated_land", 8: "seasonal_crop_land"}
    for row in range(4, ws.max_row + 1):
        district = clean(ws.cell(row, 2).value)
        if district not in provinces:
            continue
        for col, indicator in mapping.items():
            value = ws.cell(row, col).value
            if numeric(value):
                value *= 1000  # Source table is labelled (,000 Ha).
            append_record(records, year=2026, season=season, province=provinces[district], district=district, level="district", indicator=indicator, value=value, table=table)


def parse_practice_area(ws, season, provinces, records, table):
    mapping = {3: "modern_irrigated_area", 4: "erosion_control_area", 5: "agroforestry_area", 6: "inorganic_fertilizer_area", 7: "organic_fertilizer_area"}
    for row in range(6, ws.max_row + 1):
        district = clean(ws.cell(row, 2).value)
        if district not in provinces:
            continue
        for col, indicator in mapping.items():
            append_record(records, year=2026, season=season, province=provinces[district], district=district, level="district", indicator=indicator, value=ws.cell(row, col).value, table=table)


def parse_practice_percent(ws, season, provinces, records, table):
    mapping = {3: "erosion_control_farmers", 6: "mechanized_farmers", 9: "irrigating_farmers", 12: "agroforestry_farmers"}
    for row in range(5, ws.max_row + 1):
        district = clean(ws.cell(row, 2).value)
        if district not in provinces:
            continue
        for col, indicator in mapping.items():
            append_record(records, year=2026, season=season, province=provinces[district], district=district, level="district", indicator=indicator, value=ws.cell(row, col).value, table=table)


def main():
    provinces = province_map()
    records = []
    layouts = {
        "A": {"land": "Table 9", "practice_area": "Table 10", "cultivated": "Table 11", "yield": "Table 13", "production": "Table 15", "practice_pct": "Table 31", "production_header": 4, "production_data": 5},
        "B": {"land": "Table 7", "practice_area": "Table 8", "cultivated": "Table 9", "yield": "Table 11", "production": "Table 13", "practice_pct": "Table 29", "production_header": 3, "production_data": 4},
    }
    sources = {}
    for season, path in FILES.items():
        wb = load_workbook(path, read_only=True, data_only=True)
        layout = layouts[season]
        parse_national(wb["Table 6"], season, records)
        parse_crop_table(wb[layout["cultivated"]], season, "cultivated_area", 3, 4, provinces, records, layout["cultivated"])
        parse_crop_table(wb[layout["yield"]], season, "yield", 3, 4, provinces, records, layout["yield"])
        parse_crop_table(wb[layout["production"]], season, "production", layout["production_header"], layout["production_data"], provinces, records, layout["production"])
        parse_land(wb[layout["land"]], season, provinces, records, layout["land"])
        parse_practice_area(wb[layout["practice_area"]], season, provinces, records, layout["practice_area"])
        parse_practice_percent(wb[layout["practice_pct"]], season, provinces, records, layout["practice_pct"])
        sources[season] = {**SOURCE_META[season], "sha256": workbook_hash(path), "file": path.name}

    payload = {
        "dataset_version": "nisr-sas-district-2026-a-b-v1",
        "curation_method": "Direct import of published NISR aggregate Excel tables; blank cells remain missing. District yield was converted from kg/ha to metric tonnes/ha by dividing by 1,000. Land-use tables labelled thousands of hectares were converted to hectares by multiplying by 1,000.",
        "boundary_source": {
            "title": "NISR District Boundary 2022",
            "url": "https://gis-server.statistics.gov.rw/server/rest/services/Hosted/District_Boundary_2022/FeatureServer/0",
            "note": "Official NISR 2022 Census mapping boundary, simplified to 0.001 degrees for interactive display.",
        },
        "indicators": INDICATORS,
        "sources": sources,
        "records": records,
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"Wrote {len(records)} aggregate observations to {OUTPUT}")


if __name__ == "__main__":
    main()
